
import streamlit as st
import openpyxl, math, re, ast, io, os, json, uuid
from datetime import date, datetime, timezone
from docx import Document
from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_CENTER
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak
from reportlab.lib.units import mm

# -----------------------------
# APP CONFIGURATION
# -----------------------------
st.set_page_config(
    page_title="CAP776 Minor Project Evaluator",
    page_icon="🧐",
    layout="wide",
)

START = date(2026, 8, 17)
END = date(2026, 9, 21)

FEEL = {'Excellent': 5, 'Good': 4, 'Neutral': 3, 'Low': 2, 'Stressed': 1}
SAT = {'Very Satisfied': 5, 'Satisfied': 4, 'Neutral': 3, 'Unsatisfied': 2, 'Very Unsatisfied': 1}
ENERGY = {'High': 3, 'Medium': 2, 'Low': 1}
WEIGHTS = {'TPI': .15, 'AAI': .20, 'PhAI': .15, 'SRI': .20, 'TUI': .15, 'EI': .10, 'DCI': .05}

RUBRIC = [
    ('Data File & Data Continuity', 10),
    ('File Handling & Data Reading', 10),
    ('Data Validation & Exception Handling', 10),
    ('Python Fundamentals', 10),
    ('Functions & Modular Programming', 15),
    ('Python Data Structures', 10),
    ('Calculation of Activity Indices', 15),
    ('Relationship Analysis', 10),
    ('Interpretation & Findings', 5),
    ('Documentation & Reproducibility', 5),
]

REQ = [
    'Date', 'Sleep (min)', 'Fitness (min)', 'Study (min)', 'Coding (min)',
    'Class (min)', 'Classes Attended', 'Other Activities (min)',
    'Total Tracked (min)', 'Free/Unaccounted (min)', "Day's Feeling",
    'Satisfaction Level', 'Energy Level', 'Notes'
]

# -----------------------------
# CORE EVALUATOR
# -----------------------------
def read_xlsx(f):
    filename = getattr(f, 'name', '') or ''
    stem = re.sub(r'\.xlsx$', '', filename, flags=re.I)
    m = re.search(r'(\d+)', stem)
    filename_reg = m.group(1) if m else ''

    f.seek(0)
    wb = openpyxl.load_workbook(f, data_only=True)
    ws = wb['Daily Log'] if 'Daily Log' in wb.sheetnames else wb[wb.sheetnames[0]]

    headers = {}
    for c in ws[5]:
        if c.value is not None:
            h = re.sub(r'\s+', ' ', str(c.value)).strip()
            headers[h] = c.column

    missing = [x for x in REQ if x not in headers]
    if missing:
        raise ValueError('Missing columns: ' + ', '.join(missing))

    records = []
    for r in range(6, ws.max_row + 1):
        d = ws.cell(r, headers['Date']).value
        if d in (None, ''):
            continue
        records.append({
            'row': r,
            'date': d,
            'sleep': ws.cell(r, headers['Sleep (min)']).value,
            'fitness': ws.cell(r, headers['Fitness (min)']).value,
            'study': ws.cell(r, headers['Study (min)']).value,
            'coding': ws.cell(r, headers['Coding (min)']).value,
            'class': ws.cell(r, headers['Class (min)']).value,
            'classes': ws.cell(r, headers['Classes Attended']).value,
            'other': ws.cell(r, headers['Other Activities (min)']).value,
            'tracked': ws.cell(r, headers['Total Tracked (min)']).value,
            'free': ws.cell(r, headers['Free/Unaccounted (min)']).value,
            'feeling': ws.cell(r, headers["Day's Feeling"]).value,
            'satisfaction': ws.cell(r, headers['Satisfaction Level']).value,
            'energy': ws.cell(r, headers['Energy Level']).value,
            'notes': ws.cell(r, headers['Notes']).value,
        })

    workbook_reg = str(ws['F2'].value or '').strip()
    registration = filename_reg or workbook_reg or 'Not supplied'
    return {
        'name': str(ws['B2'].value or '').strip(),
        'reg': registration,
        'workbook_reg': workbook_reg,
        'filename_reg': filename_reg,
        'filename': filename,
        'section': str(ws['B3'].value or '').strip(),
        'records': records,
    }


def as_date(v):
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if isinstance(v, str):
        for fmt in ('%Y-%m-%d', '%d-%m-%Y', '%d/%m/%Y', '%d-%b-%Y'):
            try:
                return datetime.strptime(v.strip(), fmt).date()
            except Exception:
                pass
    return None


def num(v):
    return isinstance(v, (int, float)) and not isinstance(v, bool) and math.isfinite(float(v))


def validate(records):
    valid, issues, seen = [], [], set()
    numeric = ['sleep', 'fitness', 'study', 'coding', 'class', 'classes', 'other', 'tracked', 'free']

    for r in records:
        r['date'] = as_date(r['date'])
        if not r['date']:
            issues.append(f"Row {r['row']}: invalid date")
            continue
        if r['date'] in seen:
            issues.append(f"Row {r['row']}: duplicate date {r['date']}")
            continue
        seen.add(r['date'])

        activity_fields = ['sleep', 'fitness', 'study', 'coding', 'class', 'other']
        if not num(r['tracked']):
            if all(num(r[x]) for x in activity_fields):
                r['tracked'] = sum(float(r[x]) for x in activity_fields)
        if not num(r['free']) and num(r['tracked']):
            r['free'] = 1440 - r['tracked']

        bad = [x for x in numeric if not num(r[x]) or r[x] < 0]
        if bad:
            issues.append(f"Row {r['row']} ({r['date']}): invalid {', '.join(bad)}")
            continue

        if abs((r['tracked'] + r['free']) - 1440) > 0.01:
            issues.append(
                f"Row {r['row']} ({r['date']}): tracked+free is "
                f"{r['tracked'] + r['free']}, not 1440"
            )

        if r['feeling'] not in FEEL or r['satisfaction'] not in SAT or r['energy'] not in ENERGY:
            issues.append(f"Row {r['row']} ({r['date']}): invalid experience category")
            continue

        valid.append(r)

    return valid, issues


def avg(v):
    return math.fsum(v) / len(v) if v else 0


def indices(rs):
    rs_window = [r for r in rs if START <= r['date'] <= END]
    if not rs_window:
        return {k: 0 for k in ['TPI', 'AAI', 'PhAI', 'SRI', 'ABI', 'TUI', 'EI', 'DCI', 'PAI']}

    expected = (END - START).days + 1
    dci_dates = {r['date'] for r in rs_window}
    ei = avg([
        (FEEL[r['feeling']] + SAT[r['satisfaction']] + ENERGY[r['energy']]) / 3
        for r in rs_window
    ])

    x = {
        'TPI': avg([r['coding'] for r in rs_window]),
        'AAI': avg([r['study'] + r['class'] for r in rs_window]),
        'PhAI': avg([r['fitness'] for r in rs_window]),
        'SRI': avg([r['sleep'] for r in rs_window]),
        'ABI': avg([r['free'] for r in rs_window]),
        'TUI': avg([r['tracked'] for r in rs_window]),
        'EI': ei,
        'DCI': len(dci_dates) / expected * 100,
    }
    x['PAI'] = sum(WEIGHTS[k] * x[k] for k in WEIGHTS)
    return x


def corr(x, y):
    if len(x) < 2:
        return None
    mx, my = avg(x), avg(y)
    a = math.fsum((i - mx) * (j - my) for i, j in zip(x, y))
    b = math.sqrt(
        math.fsum((i - mx) ** 2 for i in x) *
        math.fsum((j - my) ** 2 for j in y)
    )
    return a / b if b else None


def relationships(rs):
    pairs = {
        'Sleep ↔ Energy': (
            [r['sleep'] for r in rs],
            [ENERGY[r['energy']] for r in rs]
        ),
        'Study ↔ Satisfaction': (
            [r['study'] for r in rs],
            [SAT[r['satisfaction']] for r in rs]
        ),
        'Coding ↔ Energy': (
            [r['coding'] for r in rs],
            [ENERGY[r['energy']] for r in rs]
        ),
    }
    return {k: corr(*v) for k, v in pairs.items()}


def report_data(f):
    f.seek(0)
    doc = Document(f)
    paras = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    tables = []
    for t in doc.tables:
        tables.append([[c.text.strip() for c in row.cells] for row in t.rows])

    found = {}
    for t in tables:
        for row in t:
            joined = ' | '.join(row)
            for k in ['PAI', 'TPI', 'AAI', 'PhAI', 'SRI', 'ABI', 'TUI', 'EI', 'DCI']:
                if k in joined:
                    nums = re.findall(r'[-+]?\d+(?:\.\d+)?', row[-1] if row else '')
                    if nums:
                        found[k] = float(nums[0])

    text = '\n'.join(paras) + '\n' + '\n'.join(' '.join(r) for t in tables for r in t)
    return found, text


def truth(reported, calc):
    out = []
    for k in ['PAI', 'TPI', 'AAI', 'PhAI', 'SRI', 'ABI', 'TUI', 'EI', 'DCI']:
        if k not in reported:
            out.append((k, None, calc[k], 'MISSING'))
        else:
            tol = .05 if k == 'EI' else .5
            out.append((
                k, reported[k], calc[k],
                'MATCH' if abs(reported[k] - calc[k]) <= tol else 'MISMATCH'
            ))
    return out


def inference_review(text, rs, rels):
    chunks = [x.strip() for x in re.split(r'(?<=[.!?])\s+', text) if len(x.split()) >= 8]
    results = []
    for s in chunks:
        low = s.lower()
        ev = []
        if 'sleep' in low:
            ev.append(f"Average sleep: {avg([r['sleep'] for r in rs]):.1f} min/day")
        if 'coding' in low:
            ev.append(f"Average coding: {avg([r['coding'] for r in rs]):.1f} min/day")
        if 'study' in low:
            ev.append(f"Average study: {avg([r['study'] for r in rs]):.1f} min/day")
        for label, r in rels.items():
            a, b = [z.strip().lower() for z in label.split('↔')]
            if a in low and b in low and r is not None:
                ev.append(f"{label}: r={r:.2f}")
        if ev:
            results.append((
                s, ev,
                bool(re.search(r'\b(cause|causes|caused|leads to|results in|improves|reduces)\b', low))
            ))
    return results


def inspect_code(f):
    f.seek(0)
    src = f.read().decode('utf-8', 'replace')
    tree = ast.parse(src)
    return {
        'src': src,
        'functions': [n.name for n in ast.walk(tree) if isinstance(n, ast.FunctionDef)],
        'imports': [
            n.name if isinstance(n, ast.Import) else n.module
            for n in ast.walk(tree) if isinstance(n, (ast.Import, ast.ImportFrom))
        ],
        'try': any(isinstance(n, ast.Try) for n in ast.walk(tree)),
        'loops': any(isinstance(n, (ast.For, ast.While)) for n in ast.walk(tree)),
        'ifs': any(isinstance(n, ast.If) for n in ast.walk(tree)),
        'lists': any(isinstance(n, ast.List) for n in ast.walk(tree)),
        'dicts': any(isinstance(n, ast.Dict) for n in ast.walk(tree)),
        'numpy': bool(re.search(r'\b(numpy|np)\b', src, re.I)),
        'pandas': bool(re.search(r'\b(pandas|pd)\b', src, re.I)),
        'classes': any(isinstance(n, ast.ClassDef) for n in ast.walk(tree)),
    }


def score(code, truths, rels, infs, issues, valid_records):
    s = {k: 0 for k, _ in RUBRIC}
    expected = (END - START).days + 1

    dci = len({
        r['date'] for r in valid_records if START <= r['date'] <= END
    }) / expected * 100

    s['Data File & Data Continuity'] = (
        10 if dci >= 95 else 9 if dci >= 90 else
        8 if dci >= 80 else 6 if dci >= 70 else 4
    )

    matches = sum(t[3] == 'MATCH' for t in truths)
    missing = sum(t[3] == 'MISSING' for t in truths)
    mismatches = sum(t[3] == 'MISMATCH' for t in truths)

    s['Calculation of Activity Indices'] = round(15 * matches / 9)

    available = sum(v is not None for v in rels.values())
    mentioned = 0
    for label in rels:
        a, b = [z.strip().lower() for z in label.split('↔')]
        if any(a in item[0].lower() and b in item[0].lower() for item in infs):
            mentioned += 1
    s['Relationship Analysis'] = min(10, 3 * available + min(1, mentioned))

    if code:
        imports = ' '.join(code['imports'])
        s['File Handling & Data Reading'] = 10 if 'openpyxl' in imports else 8
        s['Data Validation & Exception Handling'] = 10 if code['try'] else 8
        s['Python Fundamentals'] = min(
            10, 2 + 2 * code['loops'] + 2 * code['ifs'] +
            2 * bool(code['functions']) + 2 * code['lists']
        )
        s['Functions & Modular Programming'] = min(15, max(8, 3 * len(code['functions'])))
        s['Python Data Structures'] = min(10, 5 * code['lists'] + 5 * code['dicts'])
        s['Documentation & Reproducibility'] = 5 if '#' in code['src'] else 4
    else:
        s['File Handling & Data Reading'] = 8
        s['Data Validation & Exception Handling'] = 7
        s['Python Fundamentals'] = 7
        s['Functions & Modular Programming'] = 10
        s['Python Data Structures'] = 7
        s['Documentation & Reproducibility'] = 5

    if mismatches == 0 and missing == 0:
        s['Interpretation & Findings'] = 5 if infs else 4
    else:
        s['Interpretation & Findings'] = max(2, min(5, 3 + len(infs) // 2))

    return s


# -----------------------------
# REPORT PDF
# -----------------------------
def pdf_report(result):
    b = io.BytesIO()
    doc = SimpleDocTemplate(
        b, pagesize=A4,
        leftMargin=15*mm, rightMargin=15*mm,
        topMargin=15*mm, bottomMargin=15*mm
    )
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name='C', parent=styles['Title'], alignment=TA_CENTER))
    story = [
        Paragraph('CAP776 – Minor Project #1', styles['C']),
        Paragraph('Agentic Evaluation Report', styles['Heading2']),
        Spacer(1, 6)
    ]

    meta = [
        ['Student', result['student']],
        ['Registration', result['reg']],
        ['Recording period', f'{START} to {END}'],
        ['Valid days', str(len(result['valid']))],
        ['Raw score', f"{result['raw']:.1f}/100"],
        ['Scaled score', f"{result['raw']*15/100:.1f}/15"],
    ]
    t = Table(meta, colWidths=[55*mm, 115*mm])
    t.setStyle(TableStyle([
        ('GRID', (0,0), (-1,-1), .4, colors.grey),
        ('BACKGROUND', (0,0), (0,-1), colors.lightgrey)
    ]))
    story += [t, Spacer(1, 8), Paragraph('1. Parameter Truthfulness', styles['Heading2'])]

    rows = [['Index', 'Reported', 'Recalculated', 'Status']]
    rows += [
        [k, 'Missing' if r is None else f'{r:.2f}', f'{c:.2f}', status]
        for k, r, c, status in result['truth']
    ]
    t = Table(rows, colWidths=[38*mm, 38*mm, 42*mm, 30*mm], repeatRows=1)
    t.setStyle(TableStyle([
        ('GRID', (0,0), (-1,-1), .4, colors.grey),
        ('BACKGROUND', (0,0), (-1,0), colors.lightgrey),
        ('FONTSIZE', (0,0), (-1,-1), 8)
    ]))
    story += [t, Spacer(1, 8), Paragraph('2. Relationships', styles['Heading2'])]

    for k, v in result['rels'].items():
        story.append(Paragraph(
            f'<b>{k}</b>: ' + ('not calculable' if v is None else f'r = {v:.2f}'),
            styles['BodyText']
        ))

    story += [Spacer(1, 6), Paragraph('3. Rubric Scores', styles['Heading2'])]
    rows = [['Component', 'Score']]
    rows += [[k, f"{result['scores'][k]:.1f}/{m}"] for k, m in RUBRIC]
    rows += [['TOTAL', f"{result['raw']:.1f}/100"]]
    t = Table(rows, colWidths=[130*mm, 35*mm], repeatRows=1)
    t.setStyle(TableStyle([
        ('GRID', (0,0), (-1,-1), .4, colors.grey),
        ('BACKGROUND', (0,0), (-1,0), colors.lightgrey),
        ('BACKGROUND', (0,-1), (-1,-1), colors.lightgrey),
        ('FONTSIZE', (0,0), (-1,-1), 8)
    ]))
    story.append(t)

    story += [PageBreak(), Paragraph('4. Suggestions for Improvement', styles['Heading2'])]
    for x in result['suggestions']:
        story += [Paragraph('• ' + x, styles['BodyText']), Spacer(1, 3)]

    story += [Spacer(1, 6), Paragraph('5. Evaluation Notes', styles['Heading2'])]
    for x in result['notes']:
        story.append(Paragraph('• ' + x, styles['BodyText']))

    doc.build(story)
    b.seek(0)
    return b


# -----------------------------
# DATABASE (REMOVED - 100% IN-MEMORY SANDBOX)
# -----------------------------
# Submissions are processed strictly in-memory. No student data or files are stored in any database.



# -----------------------------
# UI
# -----------------------------
def evaluate_submission(xlsx, report, codef):
    data = read_xlsx(xlsx)

    filename_note = ""
    xreg = data.get('filename_reg', '')
    rreg = registration_from_filename(getattr(report, 'name', ''))
    if xreg and rreg and xreg != rreg:
        filename_note = (
            f"Registration mismatch: XLSX filename suggests {xreg}, "
            f"while report filename suggests {rreg}."
        )

    valid, issues = validate(data['records'])
    if filename_note:
        issues.append(filename_note)

    if not valid:
        raise ValueError(
            "No valid records found after validation. "
            "Please use the prescribed CAP776 workbook format."
        )

    eval_records = [r for r in valid if START <= r['date'] <= END]
    if not eval_records:
        raise ValueError(
            f"No valid records fall inside the official evaluation period "
            f"{START} to {END}."
        )

    calc = indices(eval_records)
    rels = relationships(eval_records)
    reported, text = report_data(report)
    truths = truth(reported, calc)
    infs = inference_review(text, eval_records, rels)

    code = None
    code_notes = []
    if codef:
        try:
            code = inspect_code(codef)
        except Exception as e:
            code_notes.append('Python library could not be parsed: ' + str(e))

    scores = score(code, truths, rels, infs, issues, valid)
    raw = round(sum(scores.values()), 1)

    suggestions = []
    bad = [x[0] for x in truths if x[3] == 'MISMATCH']
    missing = [x[0] for x in truths if x[3] == 'MISSING']

    if bad:
        suggestions.append('Recalculate and correct: ' + ', '.join(bad) + '.')
    if missing:
        suggestions.append('Enter all required index values in the report so they can be verified.')
    if issues:
        suggestions.append(f'Address {len(issues)} data-quality issue(s), especially date/consistency errors.')
    if calc['DCI'] < 100:
        suggestions.append(f'Improve data continuity; verified DCI is {calc["DCI"]:.1f}%.')
    if not infs:
        suggestions.append('Make findings explicitly evidence-based by citing your own averages or relationship values.')
    if any(x[2] for x in infs):
        suggestions.append('Avoid causal claims; correlation shows association, not causation.')
    if not codef:
        suggestions.append(
            'The XLSX and report were evaluated successfully. '
            'Submitting the reusable Python library is recommended when programming-process evidence is required.'
        )
    if code and (code['numpy'] or code['pandas']):
        suggestions.append('NumPy/Pandas detected; the project instructions say not to use them.')
    if not suggestions:
        suggestions.append('Maintain the evidence-based approach and improve code/documentation clarity.')

    return {
        'submission_id': str(uuid.uuid4()),
        'student': data['name'] or 'Not supplied',
        'reg': data['reg'] or 'Not supplied',
        'section': data.get('section', ''),
        'valid': valid,
        'calc': calc,
        'truth': truths,
        'rels': rels,
        'infs': infs,
        'scores': scores,
        'raw': raw,
        'suggestions': suggestions,
        'notes': issues + code_notes,
    }


def registration_from_filename(filename):
    stem = re.sub(r'\.(xlsx|docx)$', '', filename or '', flags=re.I)
    m = re.search(r'(\d+)', stem)
    return m.group(1) if m else ''


# Student submission page
st.title('🧐 CAP776 — Minor Project #1')
st.caption(
    'My Data, My Story • Agentic Evaluation • Inspired by the Original Evaluation • Sandbox Mode (Zero Database Storage) • '
    f'Official evaluation period: {START.strftime("%d %b %Y")} – {END.strftime("%d %b %Y")}'
)

with st.expander('Submission instructions', expanded=True):
    st.markdown("""
    **Upload three files:**
    1. **XLSX** — mandatory; use the prescribed daily activity workbook.
    2. **DOCX report** — mandatory; filename should contain the registration number.
    3. **Python library (.py)** — optional but recommended because it provides direct evidence of programming skills.

    **Important:** The XLSX is the numerical source of truth. The report is independently checked against recalculated values.
    """)

xlsx = st.file_uploader('1. Student XLSX — mandatory', type=['xlsx'])
report = st.file_uploader('2. Submitted Report (.docx) — mandatory', type=['docx'])
codef = st.file_uploader('3. Student Python Library (.py) — optional', type=['py'])

if st.button('🚀 SUBMIT FOR EVALUATION', type='primary', disabled=not (xlsx and report)):
    try:
        with st.spinner('Evaluating your submission...'):
            result = evaluate_submission(xlsx, report, codef)


        st.success('Evaluation completed in-memory. (Sandbox mode: No data was sent to or stored in any database.)')

        c1, c2, c3, c4 = st.columns(4)
        c1.metric('Marks', f"{result['raw']:.1f}/100")
        c2.metric('Scaled', f"{result['raw'] * 15 / 100:.1f}/15")
        c3.metric('Valid Days', len(result['valid']))
        c4.metric('DCI', f"{result['calc']['DCI']:.1f}%")

        st.subheader('Parameter Verification')
        st.dataframe(
            [
                {
                    'Index': k,
                    'Reported': 'Missing' if r is None else round(r, 2),
                    'Recalculated': round(c, 2),
                    'Status': status
                }
                for k, r, c, status in result['truth']
            ],
            use_container_width=True,
            hide_index=True
        )

        st.subheader('Relationship Analysis')
        for k, v in result['rels'].items():
            st.write(f"**{k}:** " + ('not calculable' if v is None else f"r = {v:.2f}"))

        st.subheader('Rubric')
        st.dataframe(
            [
                {'Component': k, 'Score': result['scores'][k], 'Maximum': m}
                for k, m in RUBRIC
            ],
            use_container_width=True,
            hide_index=True
        )

        st.subheader('Suggestions')
        for item in result['suggestions']:
            st.write('• ' + item)

        st.download_button(
            '📄 Download My Evaluation PDF',
            data=pdf_report(result),
            file_name=f"{result['reg']}_CAP776_Minor1_Evaluation.pdf",
            mime='application/pdf'
        )

    except Exception as e:
        st.error('Submission could not be evaluated.')
        st.warning(str(e))
